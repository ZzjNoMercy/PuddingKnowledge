"""Run a bounded local Logical Dataset Processing worker rehearsal."""

from __future__ import annotations

import argparse
import asyncio
import csv
import hashlib
import json
import shutil
import sqlite3
import tempfile
from datetime import UTC, datetime
from itertools import islice
from pathlib import Path
from typing import Any

from knowledge_contracts import Correlation, Principal
from knowledge_platform.catalog import (
    SqliteCatalogQueryRepository,
    SqliteLogicalDatasetProcessingJobStore,
    SqliteStructuredAssetWriter,
)
from knowledge_platform.structured import (
    LocalStructuredFileProvider,
    LogicalDatasetAuthoringRequest,
    LogicalDatasetAuthoringService,
    LogicalDatasetProcessingJobRequest,
    LogicalDatasetProcessingService,
    LogicalDatasetProcessingWorker,
)

_DEFAULT_OUTPUT_DIR = Path("artifacts/phase0b-local-catalog")
_SPACE_ID = "space_kb_default"
_SOURCE_ID = "phase7_shadow_source"
_DATASET_ID = "phase7_shadow_logical_dataset"


def _digest(path: Path) -> str:
    return "sha256:" + hashlib.sha256(path.read_bytes()).hexdigest()


def _path_digest(path: Path) -> str:
    return "sha256:" + hashlib.sha256(str(path.expanduser().absolute()).encode()).hexdigest()


def _bounded_csv(source: Path, target: Path) -> None:
    try:
        import openpyxl
    except ImportError as error:
        raise RuntimeError("openpyxl is required for the bounded local shadow") from error
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = list(islice(sheet.iter_rows(values_only=True), 4))
    finally:
        workbook.close()
    if len(rows) < 2:
        raise ValueError("local source does not contain a header and data row")
    header = [str(value or "").strip() for value in rows[0]]
    if not header or any(not value for value in header) or len(set(header)) != len(header):
        raise ValueError("local source header is not a bounded tabular schema")
    with target.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream, lineterminator="\n")
        writer.writerow(header)
        for row in rows[1:]:
            writer.writerow(["" if value is None else str(value) for value in row])


def _insert_source_asset(catalog: Path, *, source_path: Path, profile: object) -> None:
    now = datetime.now(UTC).isoformat()
    digest = _digest(source_path)
    columns = list(getattr(profile, "columns"))
    row_count = int(getattr(profile, "row_count"))
    uri = f"knowledge://spaces/{_SPACE_ID}/structured-assets/{_SOURCE_ID}/source"
    with sqlite3.connect(catalog) as connection:
        connection.execute(
            """
            INSERT INTO knowledge_structured_assets (
                id, space_id, source_key, document_asset_id, source_type, file_name, sheet_name,
                size_bytes, modified_at, source_uri, source_reference_digest, logical_path_digest,
                profile_uri, profile_reference_digest, content_digest, profile_status, row_count,
                column_count, columns_json, reference_status, capabilities, metadata_json, created_at, updated_at
            ) VALUES (?, ?, ?, NULL, 'csv', ?, NULL, ?, NULL, ?, ?, '', ?, ?, ?, 'ready', ?, ?, ?, 'ready', ?, '{}', ?, ?)
            """,
            (
                _SOURCE_ID,
                _SPACE_ID,
                f"shadow:{_SOURCE_ID}",
                source_path.name,
                source_path.stat().st_size,
                uri,
                digest,
                uri.replace("/source", "/profile"),
                digest,
                digest,
                row_count,
                len(columns),
                json.dumps(columns, ensure_ascii=False),
                json.dumps(["table_query"]),
                now,
                now,
            ),
        )


def run_shadow(*, output_dir: Path, canonical_catalog: Path, source_path: Path) -> dict[str, Any]:
    output_dir = output_dir.expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    canonical_catalog = canonical_catalog.expanduser().absolute()
    result: dict[str, Any] = {
        "format": "agent-knowledge-platform-phase7-local-logical-dataset-worker-shadow/v1",
        "activation": "not-activated",
        "status": "PHASE7_LOGICAL_DATASET_WORKER_SHADOW_REJECTED_NO_EXPLICIT_SOURCE",
        "catalog_path_digest": _path_digest(canonical_catalog),
        "source_path_digest": _path_digest(source_path),
        "source_digest": None,
        "derived_rows": None,
        "job": None,
        "canonical_catalog_unchanged": None,
    }
    try:
        before_catalog_digest = _digest(canonical_catalog)
        source_path = source_path.expanduser().absolute()
        result["source_digest"] = _digest(source_path)
        with tempfile.TemporaryDirectory(prefix="phase7-logical-dataset-shadow-", dir=output_dir) as temp_dir:
            temp_root = Path(temp_dir)
            bounded_source = temp_root / "bounded-source.csv"
            _bounded_csv(source_path, bounded_source)
            repository_catalog = temp_root / canonical_catalog.name
            shutil.copy2(canonical_catalog, repository_catalog)
            repository = SqliteCatalogQueryRepository(repository_catalog)
            profile = LocalStructuredFileProvider(
                asset_paths={_SOURCE_ID: bounded_source},
                asset_uris={_SOURCE_ID: f"knowledge://spaces/{_SPACE_ID}/structured-assets/{_SOURCE_ID}/source"},
            ).inspect_source(path=bounded_source)
            _insert_source_asset(repository_catalog, source_path=bounded_source, profile=profile)
            repository = SqliteCatalogQueryRepository(repository_catalog)
            principal = Principal(
                subject_id="phase7-local-logical-worker",
                scopes=(
                    "knowledge.admin",
                    "knowledge.processing",
                    f"knowledge.space:{_SPACE_ID}",
                ),
            )
            authoring = LogicalDatasetAuthoringService(
                catalog=repository,
                writer=SqliteStructuredAssetWriter(repository_catalog),
            ).create(
                principal=principal,
                correlation=Correlation("phase7-logical-authoring"),
                request=LogicalDatasetAuthoringRequest(
                    dataset_id=_DATASET_ID,
                    space_id=_SPACE_ID,
                    title="Phase 7 bounded local logical dataset",
                    source_asset_ids=(_SOURCE_ID,),
                    canonical_columns=tuple(profile.columns),
                ),
            )
            if authoring.status != "ok":
                raise RuntimeError("logical dataset authoring shadow did not enqueue")
            provider = LocalStructuredFileProvider(
                asset_paths={_SOURCE_ID: bounded_source},
                asset_uris={_SOURCE_ID: f"knowledge://spaces/{_SPACE_ID}/structured-assets/{_SOURCE_ID}/source"},
            )
            worker = LogicalDatasetProcessingWorker(
                service=LogicalDatasetProcessingService(
                    catalog=repository,
                    profiler=provider,
                    publisher=SqliteStructuredAssetWriter(repository_catalog),
                ),
                jobs=SqliteLogicalDatasetProcessingJobStore(database_path=repository_catalog),
            )
            request = LogicalDatasetProcessingJobRequest(
                dataset_id=_DATASET_ID,
                space_id=_SPACE_ID,
                source_paths={_SOURCE_ID: bounded_source},
                idempotency_key="phase7-local-logical-dataset-worker",
                principal=principal,
                correlation=Correlation("phase7-logical-processing"),
            )
            first = asyncio.run(worker.process(request))
            second = asyncio.run(worker.process(request))
            if first != second:
                raise RuntimeError("logical dataset worker terminal replay mismatch")
            result["derived_rows"] = profile.row_count
            result["job"] = {
                "job_id": first.job_id,
                "dataset_id": first.dataset_id,
                "space_id": first.space_id,
                "resource_uri": first.resource_uri,
                "content_digest": first.content_digest,
                "row_count": first.row_count,
                "second_matches_first": first == second,
            }
            with sqlite3.connect(repository_catalog) as connection:
                row = connection.execute(
                    "SELECT status, current_step, progress, lease_owner FROM knowledge_processing_jobs WHERE id = ?",
                    (first.job_id,),
                ).fetchone()
                if row != ("succeeded", "completed", 100, None):
                    raise RuntimeError("logical dataset worker terminal state is invalid")
        result["canonical_catalog_unchanged"] = _digest(canonical_catalog) == before_catalog_digest
        if not result["canonical_catalog_unchanged"]:
            raise RuntimeError("canonical Catalog changed during logical dataset worker shadow")
        result["status"] = "PHASE7_LOGICAL_DATASET_WORKER_SHADOW_PASS_NOT_ACTIVATABLE"
    except (OSError, UnicodeError, TypeError, ValueError, LookupError, RuntimeError, sqlite3.Error):
        result["status"] = "PHASE7_LOGICAL_DATASET_WORKER_SHADOW_FAILED"
        result["error"] = "phase7 local logical dataset worker shadow failed"
    return _write_report(output_dir, result)


def _write_report(output_dir: Path, result: dict[str, Any]) -> dict[str, Any]:
    report_path = output_dir / "phase7-local-logical-dataset-worker-shadow-report.json"
    report_path.write_text(json.dumps(result, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    result["report"] = str(report_path)
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=_DEFAULT_OUTPUT_DIR)
    parser.add_argument("--catalog", type=Path, default=_DEFAULT_OUTPUT_DIR / "knowledge-platform.sqlite3")
    parser.add_argument("--file", type=Path, required=True)
    args = parser.parse_args()
    result = run_shadow(output_dir=args.output_dir, canonical_catalog=args.catalog, source_path=args.file)
    print(json.dumps({"status": result["status"], "report": result["report"]}, ensure_ascii=False))
    return 0 if str(result["status"]).endswith("NOT_ACTIVATABLE") or "REJECTED" in str(result["status"]) else 1


if __name__ == "__main__":
    raise SystemExit(main())
