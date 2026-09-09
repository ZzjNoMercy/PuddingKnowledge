"""Replay logical Authoring/Processing through independent Platform processes."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import socket
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Any

from scripts.phase8_local_platform_process_shadow import (
    _mcp_http_summary,
    _post_json,
    _query_http_summary,
    _stop,
    _wait_until_ready,
)

_ROOT = Path(__file__).resolve().parents[2]
_DEFAULT_OUTPUT_DIR = _ROOT / "artifacts/phase0b-local-catalog"
_DEFAULT_CATALOG = _DEFAULT_OUTPUT_DIR / "knowledge-platform.sqlite3"
_SPACE_ID = "space_kb_default"
_SOURCE_ID = "phase8_process_source"
_DATASET_ID = "phase8_process_logical_dataset"
_IDEMPOTENCY_KEY = "phase8-local-platform-process-logical-publish"


def _digest(path: Path) -> str:
    return "sha256:" + hashlib.sha256(path.read_bytes()).hexdigest()


def _path_digest(path: Path) -> str:
    return "sha256:" + hashlib.sha256(str(path.expanduser().absolute()).encode()).hexdigest()


def _free_loopback_port() -> int:
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as listener:
        listener.bind(("127.0.0.1", 0))
        return int(listener.getsockname()[1])


def _columns(source: Path) -> list[str]:
    try:
        import openpyxl
    except ImportError as error:
        raise RuntimeError("openpyxl is required for the local logical shadow") from error
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        row = next(workbook[workbook.sheetnames[0]].iter_rows(values_only=True))
    finally:
        workbook.close()
    columns = [str(value or "").strip() for value in row]
    if not columns or any(not value for value in columns) or len(columns) != len(set(columns)):
        raise ValueError("local source header is invalid")
    return columns


def _admin_summary(status_code: int, payload: dict[str, Any]) -> dict[str, Any]:
    summary: dict[str, Any] = {"http_status": status_code, "status": payload.get("status")}
    error = payload.get("error")
    if isinstance(error, dict):
        summary["error"] = {"code": error.get("code")}
        return summary
    data = payload.get("data")
    if not isinstance(data, dict):
        return summary
    dataset = data.get("dataset")
    job = data.get("job")
    if isinstance(dataset, dict):
        summary["dataset"] = {
            key: dataset.get(key)
            for key in ("id", "space_id", "reference_status", "row_count", "content_digest")
            if key in dataset
        }
    if isinstance(job, dict):
        summary["job"] = {
            key: job.get(key) for key in ("id", "status", "current_step", "progress") if key in job
        }
    summary["evidence_count"] = len(payload.get("evidence", [])) if isinstance(payload.get("evidence"), list) else 0
    return summary


def _start_server(*, catalog: Path, source_file: Path, temp_root: Path) -> tuple[subprocess.Popen[bytes], int, Path, dict[str, Any]]:
    temp_root.mkdir(parents=True, exist_ok=True)
    ready_file = temp_root / "ready.json"
    port = _free_loopback_port()
    process = subprocess.Popen(
        [
            sys.executable,
            str(_ROOT / "backend/scripts/phase8_local_platform_process_server.py"),
            "--catalog",
            str(catalog),
            "--wiki-root",
            str(_ROOT / "docs"),
            "--temp-dir",
            str(temp_root / "server-data"),
            "--port",
            str(port),
            "--ready-file",
            str(ready_file),
            "--logical-processing-mode",
            "--logical-source-file",
            str(source_file),
        ],
        cwd=str(_ROOT),
        env={"PATH": os.environ.get("PATH", ""), "PYTHONPATH": str(_ROOT / "backend")},
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    ready = _wait_until_ready(process, ready_file, port, timeout_seconds=60)
    return process, port, temp_root / "server-data" / "knowledge-platform.sqlite3", ready


def run_shadow(
    *,
    source_file: Path,
    catalog: Path = _DEFAULT_CATALOG,
    output_dir: Path = _DEFAULT_OUTPUT_DIR,
) -> dict[str, Any]:
    source_file = source_file.expanduser().absolute()
    catalog = catalog.expanduser().absolute()
    output_dir = output_dir.expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    result: dict[str, Any] = {
        "format": "agent-knowledge-platform-phase8-local-processing-authoring-platform-process-shadow/v1",
        "activation": "not-activated",
        "status": "PHASE8_LOCAL_PROCESSING_AUTHORING_PLATFORM_PROCESS_SHADOW_FAILED",
        "transport": {"host": "127.0.0.1", "independent_processes": 2},
        "source_path_digest": _path_digest(source_file),
        "source_digest": None,
        "server": {"first": None, "second": None},
        "admin": {"authoring": None, "first_publish": None, "second_publish": None},
        "query_plane": {"first": None, "second": None},
        "restart_replay": False,
        "canonical_catalog_unchanged": None,
    }
    first_process: subprocess.Popen[bytes] | None = None
    second_process: subprocess.Popen[bytes] | None = None
    canonical_before = _digest(catalog)
    try:
        if not source_file.is_file():
            raise ValueError("explicit local source file is required")
        result["source_digest"] = _digest(source_file)
        canonical_columns = _columns(source_file)
        temporary_kwargs = {"dir": "/private/tmp"} if Path("/private/tmp").is_dir() else {}
        with tempfile.TemporaryDirectory(
            prefix="phase8-local-processing-authoring-platform-process-", **temporary_kwargs
        ) as temp_dir:
            root = Path(temp_dir)
            first_process, first_port, first_catalog, first_ready = _start_server(
                catalog=catalog, source_file=source_file, temp_root=root / "first"
            )
            create_status, create_payload = _post_json(
                first_port,
                "/v1/datasets",
                {
                    "dataset_id": _DATASET_ID,
                    "space_id": _SPACE_ID,
                    "title": "Phase 8 local platform logical dataset",
                    "source_asset_ids": [_SOURCE_ID],
                    "canonical_columns": canonical_columns,
                },
                timeout_seconds=60,
            )
            publish_body = {"space_id": _SPACE_ID, "idempotency_key": _IDEMPOTENCY_KEY}
            publish_status, publish_payload = _post_json(
                first_port,
                f"/v1/datasets/{_DATASET_ID}:publish",
                publish_body,
                timeout_seconds=60,
            )
            query_body = {
                "query": canonical_columns[0],
                "space_id": _SPACE_ID,
                "collection_id": "dataset_kb_default",
                "capability_hint": "table_query",
                "limit": 1,
            }
            query_status, query_payload = _post_json(
                first_port, "/v1/knowledge/query", query_body, timeout_seconds=60
            )
            mcp_status, mcp_payload = _post_json(
                first_port,
                "/mcp",
                {
                    "jsonrpc": "2.0",
                    "id": "logical-process-mcp-1",
                    "method": "tools/call",
                    "params": {"name": "knowledge_query", "arguments": query_body},
                },
                timeout_seconds=60,
            )
            result["server"]["first"] = {
                "capability": first_ready.get("capability"),
                "binding_present": first_ready.get("binding_present"),
                "deployment_revision": first_ready.get("deployment_revision"),
            }
            result["admin"]["authoring"] = _admin_summary(create_status, create_payload)
            result["admin"]["first_publish"] = _admin_summary(publish_status, publish_payload)
            result["query_plane"]["first"] = {
                "rest": _query_http_summary(query_status, query_payload),
                "mcp": _mcp_http_summary(mcp_status, mcp_payload),
            }
            result["first_server_shutdown_clean"] = _stop(first_process)
            if not result["first_server_shutdown_clean"]:
                raise RuntimeError("first logical sidecar did not stop cleanly")
            first_process = None
            second_process, second_port, _second_catalog, second_ready = _start_server(
                catalog=first_catalog, source_file=source_file, temp_root=root / "second"
            )
            replay_status, replay_payload = _post_json(
                second_port,
                f"/v1/datasets/{_DATASET_ID}:publish",
                publish_body,
                timeout_seconds=60,
            )
            second_query_status, second_query_payload = _post_json(
                second_port, "/v1/knowledge/query", query_body, timeout_seconds=60
            )
            second_mcp_status, second_mcp_payload = _post_json(
                second_port,
                "/mcp",
                {
                    "jsonrpc": "2.0",
                    "id": "logical-process-mcp-2",
                    "method": "tools/call",
                    "params": {"name": "knowledge_query", "arguments": query_body},
                },
                timeout_seconds=60,
            )
            result["server"]["second"] = {
                "capability": second_ready.get("capability"),
                "binding_present": second_ready.get("binding_present"),
                "deployment_revision": second_ready.get("deployment_revision"),
            }
            result["admin"]["second_publish"] = _admin_summary(replay_status, replay_payload)
            result["query_plane"]["second"] = {
                "rest": _query_http_summary(second_query_status, second_query_payload),
                "mcp": _mcp_http_summary(second_mcp_status, second_mcp_payload),
            }
            first_job = (publish_payload.get("data") or {}).get("job", {})
            second_job = (replay_payload.get("data") or {}).get("job", {})
            result["restart_replay"] = (
                create_payload.get("status") == "ok"
                and publish_payload.get("status") == "ok"
                and replay_payload.get("status") == "ok"
                and first_job.get("id") == second_job.get("id")
                and first_job.get("status") == second_job.get("status") == "succeeded"
                and first_job.get("current_step") == second_job.get("current_step") == "completed"
                and first_job.get("progress") == second_job.get("progress") == 100
                and (result["query_plane"]["first"]["rest"].get("status") == "ok")
                and (result["query_plane"]["first"]["mcp"].get("status") == "ok")
                and (result["query_plane"]["second"]["rest"].get("status") == "ok")
                and (result["query_plane"]["second"]["mcp"].get("status") == "ok")
            )
    except (OSError, RuntimeError, TypeError, ValueError, KeyError, subprocess.SubprocessError):
        result["error_type"] = "bounded_local_process_failure"
    finally:
        if first_process is not None:
            result["first_server_shutdown_clean"] = _stop(first_process)
        if second_process is not None:
            result["second_server_shutdown_clean"] = _stop(second_process)
        result["canonical_catalog_unchanged"] = _digest(catalog) == canonical_before
    if (
        result["restart_replay"]
        and result.get("first_server_shutdown_clean", True)
        and result.get("second_server_shutdown_clean", True)
        and result["canonical_catalog_unchanged"] is True
        and result["server"]["first"].get("binding_present") is True
        and result["server"]["second"].get("binding_present") is True
        and result["server"]["first"].get("deployment_revision") == "platform-local-process-v1"
        and result["server"]["second"].get("deployment_revision") == "platform-local-process-v1"
    ):
        result["status"] = "PHASE8_LOCAL_PROCESSING_AUTHORING_PLATFORM_PROCESS_SHADOW_PASS_NOT_ACTIVATABLE"
    report_path = output_dir / "phase8-local-processing-authoring-platform-process-shadow-report.json"
    report_path.write_text(json.dumps(result, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    result["report"] = str(report_path)
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--file", dest="source_file", type=Path, required=True)
    parser.add_argument("--catalog", type=Path, default=_DEFAULT_CATALOG)
    parser.add_argument("--output-dir", type=Path, default=_DEFAULT_OUTPUT_DIR)
    args = parser.parse_args()
    result = run_shadow(**vars(args))
    print(json.dumps({"status": result["status"], "report": result["report"]}, ensure_ascii=False))
    return 0 if str(result["status"]).endswith("NOT_ACTIVATABLE") else 1


if __name__ == "__main__":
    raise SystemExit(main())
