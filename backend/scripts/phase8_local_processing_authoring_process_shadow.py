"""Run a bounded independent-process Authoring/Processing continuity rehearsal."""

from __future__ import annotations

import argparse
import asyncio
import csv
import hashlib
import itertools
import json
import os
import sqlite3
import subprocess
import sys
import tempfile
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from knowledge_contracts import Correlation, Principal
from knowledge_platform.catalog import (
    SqliteCatalogQueryRepository,
    SqliteLogicalDatasetProcessingJobStore,
    SqliteStructuredAssetWriter,
)
from knowledge_platform.structured import (
    LocalStructuredFileProvider,
    LogicalDatasetAuthoringRequest,
    LogicalDatasetAuthoringService,
    LogicalDatasetProcessingJobRequest,
    LogicalDatasetProcessingService,
    LogicalDatasetProcessingWorker,
)

_DEFAULT_OUTPUT_DIR = Path("artifacts/phase0b-local-catalog")
_SPACE_ID = "space_kb_default"
_SOURCE_ID = "phase8_process_source"
_DATASET_ID = "phase8_process_logical_dataset"
_IDEMPOTENCY_KEY = "phase8-local-processing-authoring-process"
_SOURCE_URI = f"knowledge://spaces/{_SPACE_ID}/structured-assets/{_SOURCE_ID}/source"


def _digest(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        while chunk := stream.read(1024 * 1024):
            digest.update(chunk)
    return "sha256:" + digest.hexdigest()


def _path_digest(path: Path) -> str:
    return "sha256:" + hashlib.sha256(str(path.expanduser().absolute()).encode()).hexdigest()


def _bounded_csv(source: Path, target: Path) -> None:
    try:
        import openpyxl
    except ImportError as error:
        raise RuntimeError("openpyxl is required for the bounded local shadow") from error
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = list(itertools.islice(sheet.iter_rows(values_only=True), 4))
    finally:
        workbook.close()
    if len(rows) < 2:
        raise ValueError("local source does not contain a header and data row")
    header = [str(value or "").strip() for value in rows[0]]
    if not header or any(not value for value in header) or len(set(header)) != len(header):
        raise ValueError("local source header is not a bounded tabular schema")
    with target.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream, lineterminator="\n")
        writer.writerow(header)
        for row in rows[1:]:
            writer.writerow(["" if value is None else str(value) for value in row])


def _insert_pending_source(catalog: Path, *, source_path: Path, profile: object) -> None:
    now = datetime.now(UTC).isoformat()
    columns = list(getattr(profile, "columns"))
    with sqlite3.connect(catalog) as connection:
        connection.execute(
            """
            INSERT INTO knowledge_structured_assets (
                id, space_id, source_key, document_asset_id, source_type, file_name, sheet_name,
                size_bytes, modified_at, source_uri, source_reference_digest, logical_path_digest,
                profile_uri, profile_reference_digest, content_digest, profile_status, row_count,
                column_count, columns_json, reference_status, capabilities, metadata_json, created_at, updated_at
            ) VALUES (?, ?, ?, NULL, 'local', ?, NULL, ?, NULL, ?, '', '', ?, '', ?, 'missing', NULL, ?, ?, 'pending', ?, '{}', ?, ?)
            """,
            (
                _SOURCE_ID,
                _SPACE_ID,
                f"shadow:{_SOURCE_ID}",
                source_path.name,
                source_path.stat().st_size,
                _SOURCE_URI,
                _SOURCE_URI.replace("/source", "/profile"),
                str(getattr(profile, "content_digest")),
                len(columns),
                json.dumps(columns, ensure_ascii=False),
                json.dumps(["table_query"]),
                now,
                now,
            ),
        )


def _child_result(*, catalog: Path, source_path: Path, phase: str) -> dict[str, Any]:
    repository = SqliteCatalogQueryRepository(catalog)
    provider = LocalStructuredFileProvider(
        asset_paths={_SOURCE_ID: source_path},
        asset_uris={_SOURCE_ID: _SOURCE_URI},
    )
    profile = provider.inspect_source(path=source_path)
    writer = SqliteStructuredAssetWriter(catalog)
    source = repository.get_structured_asset(asset_id=_SOURCE_ID)
    if source is None:
        _insert_pending_source(catalog, source_path=source_path, profile=profile)
        repository = SqliteCatalogQueryRepository(catalog)
        source = repository.get_structured_asset(asset_id=_SOURCE_ID)
    if source is None:
        raise RuntimeError("source asset was not staged")
    if str(source.get("content_digest") or "") != profile.content_digest:
        raise RuntimeError("source digest does not match the explicit local file")
    if str(source.get("reference_status") or "") == "pending":
        writer.bind_source_asset(
            principal=Principal(
                subject_id="phase8-local-processing-authoring-process",
                scopes=("knowledge.admin", "knowledge.processing", f"knowledge.space:{_SPACE_ID}"),
            ),
            asset_id=_SOURCE_ID,
            space_id=_SPACE_ID,
            expected_content_digest=profile.content_digest,
            profile=profile,
        )
        repository = SqliteCatalogQueryRepository(catalog)
        source = repository.get_structured_asset(asset_id=_SOURCE_ID)
    if str(source.get("reference_status") or "") not in {"ready", "verified", "active"}:
        raise RuntimeError("source asset is not ready")

    principal = Principal(
        subject_id="phase8-local-processing-authoring-process",
        scopes=("knowledge.admin", "knowledge.processing", f"knowledge.space:{_SPACE_ID}"),
    )
    dataset = repository.get_structured_asset(asset_id=_DATASET_ID)
    authoring_status = "existing"
    if dataset is None:
        authoring = LogicalDatasetAuthoringService(catalog=repository, writer=writer).create(
            principal=principal,
            correlation=Correlation("phase8-process-authoring"),
            request=LogicalDatasetAuthoringRequest(
                dataset_id=_DATASET_ID,
                space_id=_SPACE_ID,
                title="Phase 8 independent-process local logical dataset",
                source_asset_ids=(_SOURCE_ID,),
                canonical_columns=tuple(profile.columns),
            ),
        )
        if authoring.status != "ok":
            raise RuntimeError("authoring did not stage pending dataset")
        authoring_status = "created"
        repository = SqliteCatalogQueryRepository(catalog)
    elif str(dataset.get("reference_status") or "") not in {"pending", "ready", "verified", "active"}:
        raise RuntimeError("logical dataset has an invalid status")

    key_digest = "sha256:" + hashlib.sha256(_IDEMPOTENCY_KEY.encode()).hexdigest()
    job_id = "logical_process_" + key_digest.removeprefix("sha256:")[:48]
    with sqlite3.connect(catalog) as connection:
        durable_before = connection.execute(
            "SELECT status, current_step, progress FROM knowledge_processing_jobs WHERE id = ?",
            (job_id,),
        ).fetchone()
    if phase == "replay" and durable_before != ("succeeded", "completed", 100):
        raise RuntimeError("durable succeeded job was not present before replay")

    worker = LogicalDatasetProcessingWorker(
        service=LogicalDatasetProcessingService(catalog=repository, profiler=provider, publisher=writer),
        jobs=SqliteLogicalDatasetProcessingJobStore(database_path=catalog),
    )
    # The worker deliberately exposes only a bounded terminal result.  Keep the
    # service diagnostic in the local child traceback for development, never in
    # the parent report or stdout contract.
    result = asyncio.run(
        worker.process(
            LogicalDatasetProcessingJobRequest(
                dataset_id=_DATASET_ID,
                space_id=_SPACE_ID,
                source_paths={_SOURCE_ID: source_path},
                idempotency_key=_IDEMPOTENCY_KEY,
                principal=principal,
                correlation=Correlation("phase8-process-processing"),
            )
        )
    )
    with sqlite3.connect(catalog) as connection:
        durable_after = connection.execute(
            "SELECT status, current_step, progress, lease_owner FROM knowledge_processing_jobs WHERE id = ?",
            (job_id,),
        ).fetchone()
    if durable_after != ("succeeded", "completed", 100, None):
        raise RuntimeError("worker did not leave a terminal durable state")
    dataset = SqliteCatalogQueryRepository(catalog).get_structured_asset(asset_id=_DATASET_ID)
    if dataset is None or str(dataset.get("reference_status") or "") not in {"ready", "verified", "active"}:
        raise RuntimeError("logical dataset is not ready after Processing")
    return {
        "phase": phase,
        "authoring_status": authoring_status,
        "durable_job_present_before": durable_before == ("succeeded", "completed", 100),
        "job_id": result.job_id,
        "dataset_id": result.dataset_id,
        "row_count": result.row_count,
        "content_digest": result.content_digest,
        "terminal_state": list(durable_after[:3]),
    }


def _run_child(args: argparse.Namespace) -> int:
    try:
        result = _child_result(catalog=args.catalog.expanduser().absolute(), source_path=args.file.expanduser().absolute(), phase=args.phase)
        print(json.dumps({"status": "ok", "result": result}, ensure_ascii=False))
        return 0
    except (OSError, sqlite3.Error, TypeError, ValueError, LookupError, RuntimeError):
        print(json.dumps({"status": "failed"}, ensure_ascii=False))
        return 1


def _invoke_child(*, catalog: Path, source_path: Path, phase: str, timeout_seconds: float) -> dict[str, Any]:
    environment = dict(os.environ)
    backend = str(Path(__file__).resolve().parents[1])
    environment["PYTHONPATH"] = backend + os.pathsep + environment.get("PYTHONPATH", "")
    completed = subprocess.run(
        [
            sys.executable,
            str(Path(__file__).resolve()),
            "--child",
            "--catalog",
            str(catalog),
            "--file",
            str(source_path),
            "--phase",
            phase,
        ],
        cwd=str(Path(__file__).resolve().parents[2]),
        env=environment,
        capture_output=True,
        text=True,
        timeout=timeout_seconds,
        check=False,
    )
    if completed.returncode != 0:
        raise RuntimeError("independent child process failed")
    try:
        payload = json.loads(completed.stdout.strip().splitlines()[-1])
    except (json.JSONDecodeError, IndexError) as error:
        raise RuntimeError("independent child process returned invalid bounded output") from error
    if not isinstance(payload, dict) or payload.get("status") != "ok" or not isinstance(payload.get("result"), dict):
        raise RuntimeError("independent child process returned a failed result")
    return payload["result"]


def run_shadow(*, output_dir: Path, canonical_catalog: Path, source_path: Path) -> dict[str, Any]:
    output_dir = output_dir.expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    canonical_catalog = canonical_catalog.expanduser().absolute()
    source_path = source_path.expanduser().absolute()
    result: dict[str, Any] = {
        "format": "agent-knowledge-platform-phase8-local-processing-authoring-process-shadow/v1",
        "activation": "not-activated",
        "status": "PHASE8_PROCESS_AUTHORING_CONTINUITY_SHADOW_REJECTED_NO_EXPLICIT_LOCAL_SOURCE",
        "catalog_path_digest": _path_digest(canonical_catalog),
        "source_path_digest": _path_digest(source_path),
        "source_digest": None,
        "bounded_source_digest": None,
        "first_process": None,
        "second_process": None,
        "restart_replay": False,
        "canonical_catalog_unchanged": None,
    }
    try:
        before_catalog = _digest(canonical_catalog)
        result["source_digest"] = _digest(source_path)
        with tempfile.TemporaryDirectory(prefix="phase8-processing-authoring-process-", dir=output_dir) as temp_dir:
            temporary_catalog = Path(temp_dir) / canonical_catalog.name
            import shutil

            shutil.copy2(canonical_catalog, temporary_catalog)
            bounded_source = Path(temp_dir) / "bounded-source.csv"
            _bounded_csv(source_path, bounded_source)
            result["bounded_source_digest"] = _digest(bounded_source)
            first = _invoke_child(catalog=temporary_catalog, source_path=bounded_source, phase="authoring_and_process", timeout_seconds=120)
            second = _invoke_child(catalog=temporary_catalog, source_path=bounded_source, phase="replay", timeout_seconds=120)
            result["first_process"] = first
            result["second_process"] = second
            result["restart_replay"] = (
                first.get("authoring_status") == "created"
                and first.get("durable_job_present_before") is False
                and second.get("authoring_status") == "existing"
                and second.get("durable_job_present_before") is True
                and first.get("job_id") == second.get("job_id")
                and first.get("content_digest") == second.get("content_digest")
                and first.get("row_count") == second.get("row_count")
                and second.get("terminal_state") == ["succeeded", "completed", 100]
            )
        result["canonical_catalog_unchanged"] = _digest(canonical_catalog) == before_catalog
        if not result["restart_replay"] or not result["canonical_catalog_unchanged"]:
            raise RuntimeError("processing/authoring continuity evidence did not satisfy the read fence")
        result["status"] = "PHASE8_PROCESS_AUTHORING_CONTINUITY_SHADOW_PASS_NOT_ACTIVATABLE"
    except (OSError, subprocess.SubprocessError, TypeError, ValueError, LookupError, RuntimeError):
        result["status"] = "PHASE8_PROCESS_AUTHORING_CONTINUITY_SHADOW_FAILED"
    report_path = output_dir / "phase8-local-processing-authoring-process-shadow-report.json"
    report_path.write_text(json.dumps(result, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    result["report"] = str(report_path)
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=_DEFAULT_OUTPUT_DIR)
    parser.add_argument("--catalog", type=Path, default=_DEFAULT_OUTPUT_DIR / "knowledge-platform.sqlite3")
    parser.add_argument("--file", type=Path, required=True)
    parser.add_argument("--child", action="store_true", help=argparse.SUPPRESS)
    parser.add_argument("--phase", choices=("authoring_and_process", "replay"), default="authoring_and_process")
    args = parser.parse_args()
    if args.child:
        return _run_child(args)
    result = run_shadow(output_dir=args.output_dir, canonical_catalog=args.catalog, source_path=args.file)
    print(json.dumps({"status": result["status"], "report": result["report"]}, ensure_ascii=False))
    return 0 if str(result["status"]).endswith("NOT_ACTIVATABLE") else 1


if __name__ == "__main__":
    raise SystemExit(main())
